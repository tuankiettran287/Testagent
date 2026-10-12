from __future__ import annotations

import csv
import io
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import urlparse

import yaml
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import settings
from .knowledge import add_source


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_REPO_FILES = 800
ALLOWED_DOCUMENTS = {".txt", ".md", ".pdf", ".docx", ".csv", ".xlsx", ".json", ".yaml", ".yml"}
SOURCE_EXTENSIONS = {".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".cs", ".go", ".rb", ".php", ".sql", ".proto", ".html", ".vue", ".svelte", ".kt", ".swift", ".rs"}
SKIP_PARTS = {".git", "node_modules", "dist", "build", "coverage", ".venv", "vendor", "bin", "obj", "__pycache__"}


def parse_document(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_DOCUMENTS:
        raise ValueError(f"Unsupported document type: {suffix}")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Document is larger than 8 MB")
    if suffix == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        tables = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + tables)
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        try:
            for sheet in workbook.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    lines.append(" | ".join(str(cell) if cell is not None else "" for cell in row))
        finally:
            workbook.close()
        return "\n".join(lines)
    text = data.decode("utf-8-sig", errors="replace")
    if suffix == ".csv":
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    return text


def ingest_document(project_id: str, name: str, data: bytes) -> str:
    return add_source(project_id, "document", name, f"document:{name}", parse_document(name, data))


def summarize_code(path: str, content: str) -> str:
    symbols = []
    for match in re.finditer(r"(?m)^\s*(?:export\s+)?(?:async\s+)?(?:def|class|function|interface|type|public\s+class|func)\s+([\w]+)", content):
        symbols.append(match.group(1))
    for match in re.finditer(r"(?m)(?:@(?:app|router)\.(?:get|post|put|patch|delete)|\b(?:GET|POST|PUT|PATCH|DELETE)\s+)/?([^\s\"']*)", content):
        symbols.append(f"endpoint {match.group(1)}")
    header = f"File: {path}\nSymbols: {', '.join(symbols[:80]) or 'none detected'}\n"
    return header + content[:MAX_FILE_BYTES]


def _eligible(path: Path) -> bool:
    return path.suffix.lower() in SOURCE_EXTENSIONS and not any(part in SKIP_PARTS for part in path.parts)


def ingest_directory(project_id: str, directory: Path, origin: str) -> dict:
    imported = []
    skipped = 0
    total_bytes = 0
    for path in sorted(directory.rglob("*")):
        if not path.is_file() or not _eligible(path.relative_to(directory)):
            continue
        if not path.resolve().is_relative_to(directory.resolve()):
            skipped += 1
            continue
        if len(imported) >= MAX_REPO_FILES:
            skipped += 1
            continue
        size = path.stat().st_size
        if size > MAX_FILE_BYTES or total_bytes + size > 50 * 1024 * 1024:
            skipped += 1
            continue
        relative = path.relative_to(directory).as_posix()
        text = path.read_text(encoding="utf-8", errors="replace")
        imported.append(add_source(project_id, "code", relative, f"{origin}:{relative}", summarize_code(relative, text)))
        total_bytes += size
    return {"imported": len(imported), "skipped": skipped, "source_ids": imported}


def ingest_repository(project_id: str, location: str) -> dict:
    parsed = urlparse(location)
    if parsed.scheme:
        if parsed.scheme != "https" or parsed.hostname not in settings.repository_hosts or parsed.username or parsed.password:
            raise ValueError("Repository URL must be HTTPS on an allowed host, without credentials")
        with tempfile.TemporaryDirectory(prefix="testagent-repo-") as temp:
            try:
                result = subprocess.run(
                    ["git", "clone", "--depth", "1", "--single-branch", "--", location, temp],
                    capture_output=True, text=True, timeout=120, check=False,
                )
            except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
                raise ValueError("Git is unavailable or repository clone timed out") from exc
            if result.returncode:
                raise ValueError("Git clone failed")
            return ingest_directory(project_id, Path(temp), location)
    path = (settings.import_root / location).resolve()
    if not path.is_relative_to(settings.import_root) or not path.is_dir():
        raise ValueError("Repository path must be a directory inside IMPORT_ROOT")
    return ingest_directory(project_id, path, str(path))


def ingest_zip(project_id: str, name: str, data: bytes) -> dict:
    if len(data) > 50 * 1024 * 1024:
        raise ValueError("ZIP is larger than 50 MB")
    imported = []
    skipped = 0
    total_bytes = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for entry in archive.infolist():
            path = Path(entry.filename)
            if entry.is_dir() or path.is_absolute() or ".." in path.parts or not _eligible(path):
                continue
            if len(imported) >= MAX_REPO_FILES or entry.file_size > MAX_FILE_BYTES or total_bytes + entry.file_size > 50 * 1024 * 1024:
                skipped += 1
                continue
            relative = path.as_posix()
            content = archive.read(entry).decode("utf-8", errors="replace")
            imported.append(add_source(project_id, "code", relative, f"zip:{name}:{relative}", summarize_code(relative, content)))
            total_bytes += entry.file_size
    return {"imported": len(imported), "skipped": skipped, "source_ids": imported}


def ingest_openapi(project_id: str, name: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("OpenAPI file is larger than 8 MB")
    try:
        spec = yaml.safe_load(data)
    except yaml.YAMLError as exc:
        raise ValueError("OpenAPI YAML/JSON cannot be parsed") from exc
    if not isinstance(spec, dict) or not ("openapi" in spec or "swagger" in spec) or not isinstance(spec.get("paths"), dict):
        raise ValueError("Invalid OpenAPI or Swagger document")
    info = spec.get("info") if isinstance(spec.get("info"), dict) else {}
    lines = [f"OpenAPI: {info.get('title', name)}"]
    for path, methods in spec["paths"].items():
        if not isinstance(methods, dict):
            continue
        for method, operation in methods.items():
            if method.lower() not in {"get", "post", "put", "patch", "delete"} or not isinstance(operation, dict):
                continue
            lines.append(f"\n{method.upper()} {path}: {operation.get('summary', '')}")
            lines.append(json.dumps({key: operation.get(key) for key in ("description", "parameters", "requestBody", "responses")}, ensure_ascii=False)[:9000])
    return add_source(project_id, "openapi", name, f"openapi:{name}", "\n".join(lines))

